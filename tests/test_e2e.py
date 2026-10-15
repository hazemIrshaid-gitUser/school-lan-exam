"""اختبار e2e: أستاذ جديد من الصفر بملفات اصطناعية.
كتاب PDF (فيه أشكال بدوائر الأرقام الصفراء) -> figures_tool -> ملفات الامتحان (بدل Claude) -> exam_data.zip جنب server.py ->
check_files.bat -> start.bat -> المعلم يراجع /key و/codes -> 14 طالب بجوالات حقيقية (Edge) -> أحداث (خروج من الصفحة، انقطاع شبكة، تعطل جهاز، تمديد، رسالة)
-> إنهاء الامتحان وتسليم عن الباقين -> مراجعة الأخطاء -> التقرير وExcel والطباعة -> بداية امتحان جديد.
ملاحظة صريحة: خطوة «Claude يجهّز الأسئلة» محاكاة بسكربت (ما في LLM بالاختبار)؛ كل الباقي حقيقي.
شغّله لحاله مع تفاصيل: pytest test_e2e.py -s       للقطات: SHOTS_DIR=مجلد"""
import io, json, os, pathlib, random, re, shutil, time, urllib.parse, zipfile
from collections import Counter

import openpyxl
import pytest

pw = pytest.importorskip("playwright.sync_api")
pymupdf = pytest.importorskip("pymupdf")
pypdf = pytest.importorskip("pypdf")
from conftest import PKG_SRC, Q_HEAD, S_HEAD
from test_browser import browser, phone, url, clean, pick, submit         # noqa: F401  (fixtures)
from test_http import Server, TEACHER, free_port
from test_windows_env import Run, SYS32

SHOTS = os.environ.get("SHOTS_DIR")
TITLE = "اختبار الوحدة الثانية — الحركة (e2e)"
PASSWORD = "T3acher!x"
FINAL = 20
PICK = {"hard": 2, "medium": 4, "easy": 4}
NAMES_A = ["أحمد خالد الحسن", "سارة محمد العلي", "عمر يوسف النمر", "لينا عبد الله", "يزن فؤاد حداد", "ريم سامر القاسم", "طارق ناصر", "دانا إبراهيم"]
NAMES_B = ["مالك أيمن", "هبة وليد", "زيد حسام الدين", "نور الهدى", "كريم باسل", "جود رائد"]
ALL = [("العاشر", "أ", n) for n in NAMES_A] + [("العاشر", "ب", n) for n in NAMES_B]
LOG = []


def step(msg):
    LOG.append(msg); print(f"\n[e2e] {msg}", flush=True)


# =============================== مواد اصطناعية ===============================
def make_book(path):
    doc = pymupdf.open()
    for pno in range(1, 7):
        page = doc.new_page(width=595, height=794)
        page.insert_text((60, 40), f"Physics Grade 10 - Unit 2 - Page {pno}", fontsize=10)
        page.insert_text((60, 70), "Motion in one dimension and the geometry of figures", fontsize=11)
        if pno in (3, 4, 5):
            for k in range(2):
                y = 170 + k * 300
                page.draw_circle((90, y + 18), 10, color=None, fill=(1, 0.8, 0.2)); page.insert_text((85, y + 22), str(pno * 10 + k), fontsize=11)
                page.insert_text((120, y), f"Find the value of x in the figure number {pno}.{k} shown below", fontsize=11)
                cx, cy, r = 300 + k * 30, y + 100, 50 + pno * 2
                page.draw_circle((cx, cy), r, color=(0, 0, 0), width=1.5)
                page.draw_line((cx - r, cy), (cx + r, cy), color=(0, 0, 0), width=1.5); page.draw_line((cx, cy - r), (cx + r * 0.7, cy + r * 0.7), color=(0, 0, 0), width=1.5)
                for t, (dx, dy) in {"O": (4, -4), "A": (-r - 14, 4), "B": (r + 6, 4)}.items(): page.insert_text((cx + dx, cy + dy), t, fontsize=11)
    doc.save(str(path)); doc.close()


def make_exam_files(d, figs):
    """بدل Claude: 30 سؤال (10 صعب بعلامتين، 10 وسط، 10 سهل)، 6 منها بأشكال من الكتاب. بيرجّع QA: نص السؤال -> الصحيح والعلامة"""
    rng = random.Random(11); QA, rows, idx = {}, [], 0
    for lvl, count, marks in (("صعب", 10, 2), ("وسط", 10, 1), ("سهل", 10, 1)):
        for j in range(count):
            idx += 1
            if j < 2 and figs:                                     # شكلين لكل مستوى
                fig = figs.pop(0); ans = rng.choice([30, 45, 60, 75, 120]); opts = [f"{ans}°", f"{ans + 15}°", f"{ans - 10}°", f"{ans * 2}°"]
                text = f"من الشكل الآتي (س{idx}) أوجد قيمة x: [img:{fig}]"
            else:
                a, b = rng.randint(2, 9), rng.randint(3, 12); ans = a * b * a
                opts = [str(ans), str(ans + 10), str(ans - 4), str(ans * 2)]
                text = f"سيارة تتحرك بسرعة {a * b} م/ث لمدة {a} ث، ما المسافة المقطوعة بالأمتار؟ (س{idx})"
            pos = rng.randrange(4); opts = opts[1:pos + 1] + [opts[0]] + opts[pos + 1:]          # الصحيح بمكان عشوائي
            letter = "أبجد"[opts.index(opts[pos])]
            key = re.sub(r"\s+", " ", re.sub(r"\[img:[^\]]+\]", "", text)).strip()
            rows.append([text, *opts, letter, lvl, f"ص {30 + idx}", f"الحل: {ans}", marks, ""])
            QA[key] = {"right": re.sub(r"\s+", " ", opts[pos]).strip(), "marks": marks, "level": {"صعب": "hard", "وسط": "medium", "سهل": "easy"}[lvl]}
    wb = openpyxl.Workbook(); ws = wb.active; ws.append(Q_HEAD)
    for r in rows: ws.append(r)
    wb.save(d / "questions.xlsx")
    wb = openpyxl.Workbook(); ws = wb.active; ws.append(S_HEAD)
    manual = {0: 123456, 1: "٤٥٦٧٨٩", 2: "012345"}                  # رموز يدوية: رقم، أرقام عربية، نص بصفر بادئ
    for i, (g, s, n) in enumerate(ALL): ws.append([g, s, n, manual.get(i, "")])
    wb.save(d / "students.xlsx")
    (d / "settings.txt").write_text("\n".join([f"عنوان_الامتحان = {TITLE}", "المدة_بالدقائق = 30", "عدد_الأسئلة_الصعبة = 2", "عدد_الأسئلة_المتوسطة = 4", "عدد_الأسئلة_السهلة = 4", "عدد_الأسئلة_الكلي = 10",
                                     "خلط_الأسئلة = نعم", "خلط_الخيارات = نعم", "إظهار_العلامة_للطالب = نعم", "منع_النسخ = نعم", "مراجعة_الأخطاء = بعد_الإنهاء", "رمز_الدخول = نعم",
                                     f"كلمة_سر_المعلم = {PASSWORD}", f"العلامة_النهائية = {FINAL}"]) + "\n", encoding="utf-8")
    return QA


class _FakePkg:
    """أقل شي بيحتاجه Server للأوامر (get/post): مسار الحزمة"""
    def __init__(self, path): self.path = path
    def p(self, n): return self.path / n


# =============================== أدوات المعلم والطالب ===============================
def teacher_login(page, w):
    page.goto(url(w) + "login"); page.fill("#p", PASSWORD); page.click("#b"); page.wait_for_url(re.compile(r"/monitor"))


def norm(t): return re.sub(r"\s+", " ", t.replace("⁦", "").replace("⁩", "")).strip()


def read_cards(page):
    return page.evaluate("""()=>[...document.querySelectorAll('#qs .card')].map(c=>{const q=c.querySelector('.q').cloneNode(true);q.querySelectorAll('.mk,img').forEach(e=>e.remove());
        return {q:q.textContent,o:[...c.querySelectorAll('label.opt span')].map(s=>s.textContent)}})""")


def answer(page, QA, wrong=0, only=None, skip=()):
    """يجاوب: أول wrong سؤال غلط والباقي صح. بيرجّع قائمة (نص السؤال، صح؟) للأسئلة اللي جاوبها"""
    out, done_wrong = [], 0
    for i, c in enumerate(read_cards(page)):
        if only is not None and i not in only: continue
        if i in skip: continue
        key = norm(re.sub(r"^\d+\.\s*", "", c["q"], count=1)); right = QA[key]["right"]; opts = [norm(o) for o in c["o"]]
        ok = done_wrong >= wrong
        j = opts.index(right) if ok else next(k for k, o in enumerate(opts) if o != right)
        if not ok: done_wrong += 1
        pick(page, i, j)
        out.append((key, ok))
    return out


def expected(QA, picks, total_marks):
    """العلامة المتوقعة بعد التحويل لـ 20: مجموع علامات الصح / مجموع علامات الورقة * 20"""
    got = sum(QA[k]["marks"] for k, ok in picks if ok)
    return round(got / total_marks * FINAL, 2)


def student_login(page, w, name, section, code, grade="العاشر"):
    page.goto(url(w)); page.select_option("#g", grade); page.select_option("#sec", section); page.select_option("#nm", name)
    page.fill("#code", code); page.click("#start"); page.wait_for_selector("#s-exam:not(.hide)")


def done_ok(page, timeout=20000): page.wait_for_function("document.getElementById('status').textContent.includes('✓')", timeout=timeout)


def shot(page, name):
    if SHOTS:
        pathlib.Path(SHOTS).mkdir(parents=True, exist_ok=True); page.screenshot(path=str(pathlib.Path(SHOTS) / f"e2e_{name}.png"), full_page=True)


# =============================== السيناريو ===============================
def test_new_teacher_full_exam_day(browser, phone, tmp_path):
    t_start = time.time()
    # ---------- 1) المواد: كتاب + figures_tool ----------
    step("1) كتاب PDF اصطناعي + قص الأشكال بـ figures_tool.py")
    work = tmp_path / "مواد الأستاذ"; work.mkdir()
    book = work / "كتاب الفيزياء — الوحدة 2.pdf"; make_book(book)
    figs_dir = work / "figs"
    import subprocess, sys
    r = subprocess.run([sys.executable, "-X", "utf8", str(PKG_SRC / "figures_tool.py"), "detect", str(book), "3-5", "--out", str(figs_dir)], capture_output=True, text=True, encoding="utf-8")
    assert r.returncode == 0, r.stdout + r.stderr
    index = json.loads((figs_dir / "index.json").read_text(encoding="utf-8"))
    assert len(index) == 6 and [f["question_badge"] for f in index] == ["30", "31", "40", "41", "50", "51"], [(f["file"], f["question_badge"]) for f in index]
    fig_names = [f["file"] for f in index]

    step("2) «Claude» (محاكاة) يجهّز الأسئلة والصور -> exam_data.zip (الطلاب والإعدادات عند المشرف)")
    data_dir = work / "exam_data"; (data_dir / "images").mkdir(parents=True)
    for f in fig_names: shutil.copy(figs_dir / f, data_dir / "images" / f)
    QA = make_exam_files(data_dir, list(fig_names))
    assert len(QA) == 30
    z = work / "exam_data.zip"
    with zipfile.ZipFile(z, "w", zipfile.ZIP_DEFLATED) as zf:
        for p in sorted(data_dir.rglob("*")):
            if p.is_file() and p.name not in ("students.xlsx", "settings.txt"): zf.write(p, p.relative_to(data_dir).as_posix())

    # ---------- 2) الأستاذ ينسخ الحزمة بمسار عربي فيه مسافات وأقواس ويحط الزيب ----------
    step("3) نسخ الحزمة لمجلد بمسار صعب + وضع exam_data.zip جنب server.py")
    home = tmp_path / "امتحان الفيزياء (العاشر) & 2026"
    shutil.copytree(PKG_SRC, home, ignore=shutil.ignore_patterns("__pycache__", "*.pyc"))
    shutil.copy(z, home / "exam_data.zip")
    for n in ("students.xlsx", "settings.txt"): shutil.copy(data_dir / n, home / n)          # ملفات المشرف (قائمة المختبر وإعدادات طلب الأستاذ)

    # ---------- 3) check_files.bat ----------
    step("4) check_files.bat")
    chk = Run("check_files.bat", home); code = chk.finish(60)
    assert code == 0 and "OK - the exam files are valid" in chk.text, chk.text
    rep = (home / "check_report.txt").read_text(encoding="utf-8-sig")
    assert "بنك الأسئلة    : 30 سؤال" in rep and "بالامتحان      : 10 سؤال" in rep and "العلامة الكاملة: 20" in rep and "الطلاب         : 14" in rep
    assert "6 صورة بالأسئلة" in rep and "رمز الدخول     : مفعّل" in rep
    step("   تقرير الفحص:\n" + "\n".join("     " + l for l in rep.splitlines()[:14]))
    assert not [l for l in rep.splitlines() if "⚠" in l and "بدون حل" not in l], rep            # لا تنبيهات غير متوقعة

    # ---------- 4) start.bat ----------
    step("5) start.bat")
    port = free_port(); srv = Run("start.bat", home, env={"EXAM_PORT": str(port)})
    try:
        assert srv.wait_for("Give students this link"), srv.text
        assert "WARNING" not in srv.text.split("Teacher pages")[0]
        w = Server(_FakePkg(home)); w.port = port
        run_exam_day(browser, phone, w, home, QA, port, srv, t_start)
    finally:
        srv.kill()


def run_exam_day(browser, phone, w, home, QA, port, srv, t_start):
    # ---------- 5) المعلم يراجع المفتاح والرموز ----------
    step("6) المعلم: الدخول، مفتاح الإجابات، الرموز، معاينة")
    ctx = browser.new_context(viewport={"width": 1280, "height": 900}, locale="ar-JO")
    ctx.grant_permissions(["clipboard-read", "clipboard-write"], origin=f"http://127.0.0.1:{port}")
    T = ctx.new_page(); T.set_default_timeout(10000)
    T.dlg, T.answers = [], []
    def on_dialog(d):
        T.dlg.append((d.type, d.message))
        d.accept(T.answers.pop(0)) if (d.type == "prompt" and T.answers) else d.accept()
    T.on("dialog", on_dialog); errs = []; T.on("pageerror", lambda e: errs.append(str(e)))
    teacher_login(T, w)
    T.goto(url(w) + "key"); T.wait_for_load_state()
    assert len(T.query_selector_all("div.q")) == 30 and T.inner_text("body").count("✓") >= 30
    imgs = T.query_selector_all("img.qimg"); assert len(imgs) == 6 and all(T.evaluate("i=>i.complete&&i.naturalWidth>0", i) for i in imgs)
    T.goto(url(w) + "codes"); codes_html = T.content()
    slips = dict(re.findall(r"<div class=\"?slip\"?>(.*?)<span class=\"?code\"?>(\d+|[^<]+)</span>", codes_html))
    assert len(slips) == 14, len(slips)
    codes = {n: slips[n] for _, _, n in ALL}
    assert codes[NAMES_A[0]] == "123456" and codes[NAMES_A[1]] == "456789" and codes[NAMES_A[2]] == "012345"      # اليدوية بأشكالها
    assert all(len(c) == 6 for c in codes.values())
    T.goto(url(w) + "teacher"); T.wait_for_selector("#s-exam:not(.hide)", timeout=15000)
    assert len(T.query_selector_all("#qs .card")) == 30                                                        # معاينة المعلم: كل البنك
    shot(T, "01_teacher_preview")
    T.goto(url(w) + "monitor"); T.wait_for_selector("#cards .card")
    assert [int(x) for x in T.eval_on_selector_all("#cards .n", "e=>e.map(x=>x.textContent)")] == [0, 0, 0, 14]

    # ---------- 6) الطلاب يدخلوا ----------
    step("7) 13 طالب بجوالات: دخول بالرمز وبدء الامتحان (طالب 14 غايب)")
    phones = {}; plans = {}
    for g, s, n in ALL[:13]:
        pg = phone(); pg.set_default_timeout(15000); student_login(pg, w, n, s, codes[n]); phones[n] = (pg, s)
    shot(phones[NAMES_A[0]][0], "02_student_exam")
    papers = {n: [norm(re.sub(r"^\d+\.\s*", "", c["q"], count=1)) for c in read_cards(pg)] for n, (pg, s) in phones.items()}
    assert all(len(set(v)) == 10 for v in papers.values()) and len({tuple(v) for v in papers.values()}) > 6      # أوراق مختلفة، 10 أسئلة بدون تكرار
    totals = {n: sum(QA[k]["marks"] for k in v) for n, v in papers.items()}
    assert set(totals.values()) == {12}                                                                          # 2*2 + 4 + 4 لكل ورقة
    cnt = Counter(QA[k]["level"] for k in papers[NAMES_A[0]]); assert cnt == Counter(PICK)
    T.evaluate("load()"); T.wait_for_timeout(300)
    c0 = [int(x) for x in T.eval_on_selector_all("#cards .n", "e=>e.map(x=>x.textContent)")]
    assert c0[0] == 0 and c0[1] + c0[2] == 13 and c0[3] == 1, c0

    # رمز غلط ثم صحيح لطالب 14 (غايب): ما بيبدأ لاحقاً
    # ---------- 7) الإجابات والأحداث ----------
    step("8) طلاب 1-8: حل بأخطاء متفاوتة وتسليم")
    picks = {}
    for i, n in enumerate(NAMES_A):
        pg, s = phones[n]; picks[n] = answer(pg, QA, wrong=i % 5)
        submit(pg); done_ok(pg)
        assert f"{expected(QA, picks[n], 12):g} / {FINAL}" in clean(pg.inner_text("#score")).replace("⁦", ""), (n, expected(QA, picks[n], 12), clean(pg.inner_text("#score")))
    shot(phones[NAMES_A[0]][0], "03_student_done")

    step("9) طالب 9: يخرج من الصفحة مرتين")
    n9 = NAMES_B[0]; pg9 = phones[n9][0]
    pg9.evaluate("dispatchEvent(new Event('blur'))"); pg9.evaluate("dispatchEvent(new Event('focus'))"); pg9.evaluate("dispatchEvent(new Event('blur'))")
    assert pg9.is_visible("#away"); picks[n9] = answer(pg9, QA, wrong=2); pg9.wait_for_timeout(2200)
    submit(pg9); done_ok(pg9)

    step("10) طالب 10: ينقطع النت 6 ثواني أثناء الامتحان ويرجع")
    n10 = NAMES_B[1]; pg10 = phones[n10][0]
    picks[n10] = answer(pg10, QA, wrong=1, only=set(range(0, 5))); pg10.wait_for_timeout(2200)
    pg10.context.set_offline(True); picks[n10] += answer(pg10, QA, only=set(range(5, 10))); pg10.wait_for_timeout(6000)            # بيحل بدون نت
    pg10.context.set_offline(False); pg10.wait_for_timeout(9500)                                                                   # نبضة الـ ping كل 8 ثواني
    submit(pg10); done_ok(pg10)

    step("11) طالب 11: جهازه يتعطل: المعلم يعمل «نقل لجهاز آخر» وبيكمل من جوال ثاني بنفس الورقة والإجابات")
    n11 = NAMES_B[2]; pg11 = phones[n11][0]
    first = answer(pg11, QA, wrong=1, only=set(range(0, 6))); pg11.wait_for_timeout(2500)
    T.evaluate("load()"); T.click("text=فتح كل الشعب"); T.wait_for_timeout(300)
    T.locator("details.sec tr", has_text=n11).get_by_text("نقل لجهاز آخر", exact=True).click(); T.wait_for_timeout(600)
    assert any("نقل الطالب لجهاز آخر" in m for _, m in T.dlg)
    pg11.context.close()                                                                                        # الجهاز القديم مات
    new11 = phone(); new11.set_default_timeout(15000); student_login(new11, w, n11, "ب", codes[n11])              # لازم يدخل رمزه من جديد
    assert [norm(re.sub(r"^\d+\.\s*", "", c["q"], count=1)) for c in read_cards(new11)] == papers[n11]            # نفس الورقة بنفس الترتيب
    assert new11.eval_on_selector_all("#qs input:checked", "e=>e.length") == 6                                  # إجاباته محفوظة
    picks[n11] = first + answer(new11, QA, only=set(range(6, 10))); phones[n11] = (new11, "ب")
    new11.wait_for_timeout(2200); submit(new11); done_ok(new11)

    step("12) طالب 12: المعلم يمدّد له 5 دقائق من واجهة المراقبة")
    n12 = NAMES_B[3]; pg12 = phones[n12][0]
    before = pg12.inner_text("#clock"); T.answers.append("5")
    T.locator("details.sec tr", has_text=n12).get_by_text("+ وقت", exact=True).click(); T.wait_for_timeout(500)
    picks[n12] = answer(pg12, QA, wrong=3); pg12.wait_for_timeout(2500)                                           # تغيير الإجابة يفعّل ping فيقرأ الوقت الجديد
    mm = int(pg12.inner_text("#clock").split(":")[0]); assert mm >= 33, (before, pg12.inner_text("#clock"))
    submit(pg12); done_ok(pg12)

    step("13) رسالة من المعلم لكل الطلاب")
    n13 = NAMES_B[4]; pg13 = phones[n13][0]; T.answers.append("باقي 10 دقائق، تأكدوا من إجاباتكم")
    T.click("text=📢 رسالة للطلاب"); T.wait_for_timeout(500)
    picks[n13] = answer(pg13, QA, wrong=2, only=set(range(0, 6)))                                                    # يحل 6 من 10 (بيفعّل ping)
    pg13.wait_for_selector("#notice:not(.hide)", timeout=12000); assert "باقي 10 دقائق" in pg13.inner_text("#noticeText")
    pg13.wait_for_timeout(2500)
    step("14) طالب 13: يتعطل جواله نهائياً بدون تسليم (بعد حل 6 أسئلة)")
    pg13.context.close()                                                                                          # لا تسليم ولا رجوع

    # ---------- 8) المعلم يراقب ثم ينهي ----------
    step("15) المراقبة: 12 سلّموا، طالب 13 بدأ ولم يسلّم، طالب 14 غايب")
    T.evaluate("load()"); T.wait_for_timeout(400)
    cards = [int(x) for x in T.eval_on_selector_all("#cards .n", "e=>e.map(x=>x.textContent)")]
    assert cards[0] == 12 and cards[1] + cards[2] == 1 and cards[3] == 1, cards                                   # [سلّموا، يمتحنون، انقطعوا، لم يبدأوا]
    shot(T, "04_monitor_before_close")
    row9 = T.locator("details.sec tr", has_text=n9).inner_text(); assert "⚠" in row9 and "2 مرة" in row9
    T.click("text=⏹ إنهاء الامتحان للجميع"); T.wait_for_selector("#closedbar", state="visible")
    T.click("text=📥 تسليم من لم يسلّم"); T.wait_for_function("document.querySelectorAll('#cards .n')[0].textContent==='13'", timeout=10000)
    assert any("تسليم 1 طالب" in m for _, m in T.dlg)
    sub13 = [json.loads(l) for l in (home / "submissions.jsonl").read_text(encoding="utf-8").splitlines() if n13 in l][0]
    got13 = sum(QA[k]["marks"] for k, ok in picks[n13] if ok)
    assert sub13["auto"] == "تسليم من المعلم" and sub13["score"] == round(got13 / 12 * FINAL, 2), (sub13["score"], got13)       # الأسئلة بدون إجابة = خاطئة

    step("16) طالب 14 (غايب) يحاول يدخل بعد الإنهاء")
    late = phone(); late.set_default_timeout(15000); n14 = NAMES_B[5]
    late.goto(url(w)); late.select_option("#g", "العاشر"); late.select_option("#sec", "ب"); late.select_option("#nm", n14); late.fill("#code", codes[n14]); late.click("#start"); late.wait_for_timeout(900)
    assert any("أنهى المعلم الامتحان" in m for m in late.dialogs) and late.is_visible("#s-login")

    step("17) مراجعة الأخطاء بعد الإنهاء تظهر على شاشة الطلاب")
    pg1 = phones[NAMES_A[0]][0]                                                                                    # أخطاءه 0
    pg3 = phones[NAMES_A[2]][0]                                                                                    # أخطاءه 2
    pg3.wait_for_selector("#review .rv", timeout=16000); assert len(pg3.query_selector_all("#review .rv")) == 2
    pg1.wait_for_function("document.getElementById('review').textContent.includes('لا توجد أخطاء')", timeout=16000)
    shot(pg3, "05_student_review")

    # ---------- 9) النتائج والتقارير ----------
    step("18) تحقق من كل علامة (سجل السيرفر مقابل الحساب المستقل)")
    recs = {r["name"]: r for r in (json.loads(l) for l in (home / "submissions.jsonl").read_text(encoding="utf-8").splitlines())}
    assert set(recs) == {n for _, _, n in ALL[:13]}
    for n, pk in picks.items():
        assert abs(recs[n]["score"] - expected(QA, pk, 12)) < 0.011 and recs[n]["total"] == FINAL, (n, recs[n]["score"], expected(QA, pk, 12))
    assert recs[n9]["leave"].startswith("2 مرة") and recs[n9]["late"] == 0
    mj = w.get("/monitor.json")[1]; assert mj["c"] == {"done": 13, "live": 0, "lost": 0, "none": 1} and mj["closed"] is True

    step("19) التقرير وExcel وورقات الطلاب والطباعة")
    T.goto(url(w) + "report"); body = T.inner_text("body")
    assert "ملخص الشعب" in body and "تحليل الأسئلة" in body and not errs
    xl = T.request.get(url(w) + "report.xlsx"); assert xl.status == 200
    wb = openpyxl.load_workbook(io.BytesIO(xl.body()), data_only=True); assert wb.sheetnames[0] == "ملخص" and len(wb.sheetnames) >= 4
    summ = {r[1]: r for r in wb["ملخص"].iter_rows(min_row=2, values_only=True)}
    for sec, names in (("أ", NAMES_A), ("ب", NAMES_B[:5])):
        want = round(sum(expected(QA, picks[n], 12) / FINAL * 100 for n in names) / len(names), 1)
        assert abs(summ[sec][4] - want) < 0.15, (sec, summ[sec], want)
    qsheet = list(wb["تحليل الأسئلة"].iter_rows(min_row=2, values_only=True)); assert sum(r[3] for r in qsheet) == 13 * 10
    res_csv = list(__import__("csv").reader(io.StringIO((home / "results.csv").read_text(encoding="utf-8-sig")))); assert len(res_csv) == 14
    T.goto(url(w) + "sheet?g=" + urllib.parse.quote("العاشر") + "&s=" + urllib.parse.quote("أ")); T.wait_for_load_state()
    data = T.pdf(format="A4", print_background=True); pages = pypdf.PdfReader(io.BytesIO(data)).pages
    assert len(pages) >= 8 and T.inner_text("body").count("العلامة:") == 8                                         # ورقة لكل طالب بالشعبة أ
    shot(T, "06_sheet_section")
    T.goto(url(w) + "codes"); data = T.pdf(format="A4"); assert len(pypdf.PdfReader(io.BytesIO(data)).pages) >= 2    # قصاصات الرموز

    step("20) بداية امتحان جديد: أرشفة وتصفير")
    T.goto(url(w) + "monitor"); T.wait_for_selector("#cards .card"); T.click("text=بداية امتحان جديد"); T.wait_for_function("document.querySelectorAll('#cards .n')[3].textContent==='14'", timeout=10000)
    arch = [d for d in os.listdir(home) if d.startswith("archive_")]; assert arch and (home / arch[0] / "submissions.jsonl").exists()
    assert not (home / "submissions.jsonl").exists() or (home / "submissions.jsonl").read_text(encoding="utf-8") == ""
    assert not errs, errs
    ctx.close()
    step(f"تمّ: e2e كامل بـ {time.time() - t_start:.0f} ثانية | 14 طالب | 13 سلّموا (12 بأنفسهم + 1 من المعلم) | العلامات كلها مطابقة")

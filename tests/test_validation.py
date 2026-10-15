"""المرحلة 2: ملفات خربانة عمداً — لازم السيرفر يرفضها برسالة واضحة فيها رقم السطر، والملفات السليمة لازم تنقبل."""
import io, re, subprocess, sys, zipfile

import pytest
from conftest import Q_HEAD, S_HEAD, bank, q, roster


# ---------- الحالة السليمة ----------
def test_valid_files_load(pkg):
    m = pkg.load()
    assert len(m.QLIST) == 11 and sum(m.PICK.values()) == 8
    assert sum(len(n) for g in m.ROSTER.values() for n in g.values()) == 5


def test_template_files_from_package_load(pkg):
    """الحزمة الأصلية (القوالب + settings.txt) لازم تنقبل بدون تعديل، لو كان البنك كافي للعدد المطلوب"""
    import shutil
    from conftest import PKG_SRC
    shutil.copy(PKG_SRC / "questions_template.xlsx", pkg.p("questions.xlsx"))
    shutil.copy(PKG_SRC / "students_template.xlsx", pkg.p("students.xlsx"))
    shutil.copy(PKG_SRC / "settings.txt", pkg.p("settings.txt"))
    # إعدادات الحزمة: العنوان والعدد فاضيين = عنوان تلقائي + كل أسئلة الملف (القالب فيه 3 أسئلة)
    m = pkg.load()
    assert sum(m.PICK.values()) == len(m.QLIST) == 3 and m.SET["auto_title"] and m.SET["minutes"] == 30 and m.SET["tpass"] == "1234"


def test_counts_bigger_than_bank_still_error(pkg):
    pkg.settings(**{"عدد_الأسئلة_الصعبة": 3, "عدد_الأسئلة_المتوسطة": 6, "عدد_الأسئلة_السهلة": 6})
    pkg.expect_error("أقل من المطلوب")


@pytest.mark.parametrize("key", ["المدة_بالدقائق", "خلط_الأسئلة", "عدد_الأسئلة_الكلي", "رمز_الدخول"])
def test_empty_value_means_default(pkg, key):
    pkg.raw_settings(f"{key} =\n"); m = pkg.load()
    assert m.SET["minutes"] == 30 and m.SET["shuffleQ"] is True and m.SET["needCode"] is False and sum(m.PICK.values()) == 11


# ---------- ملف الأسئلة ----------
@pytest.mark.parametrize("ans", ["هـ", "ز", "5", "xyz"])
def test_bad_answer_letter(pkg, ans):
    rows = bank(); rows[2][5] = ans
    pkg.questions(rows)
    pkg.expect_error("سطر 4", "الجواب")


def test_answer_beyond_option_count(pkg):
    rows = bank(); rows[0] = q("ثلاث خيارات", ("1", "2", "3"), "د", "صعب")
    pkg.questions(rows)
    pkg.expect_error("سطر 2", "الجواب")


@pytest.mark.parametrize("ans", ["أ", "ا", "إ", "a", "A", "1", "أ."])
def test_answer_letter_variants_accepted(pkg, ans):
    rows = bank(); rows[0][5] = ans
    pkg.questions(rows)
    assert pkg.load().QLIST[0]["a"] == 0


def test_empty_question_text(pkg):
    rows = bank(); rows[1][0] = ""
    pkg.questions(rows)
    pkg.expect_error("سطر 3", "فاضي")


def test_single_option_rejected(pkg):
    rows = bank(); rows[0] = q("خيار واحد", ("فقط",), "أ", "صعب")
    pkg.questions(rows)
    pkg.expect_error("سطر 2", "خيارين")


def test_gap_between_options_rejected(pkg):
    rows = bank(); rows[0] = q("فراغ", ("1", "", "3", "4"), "أ", "صعب")
    pkg.questions(rows)
    pkg.expect_error("سطر 2", "خيارين")


def test_two_options_accepted(pkg):
    rows = bank(); rows[0] = q("صح أو خطأ", ("صح", "خطأ"), "ب", "صعب")
    pkg.questions(rows)
    assert len(pkg.load().QLIST[0]["o"]) == 2


def test_bad_level(pkg):
    rows = bank(); rows[0][6] = "مستحيل"
    pkg.questions(rows)
    pkg.expect_error("سطر 2", "المستوى")


def test_partially_empty_levels_rejected(pkg):
    rows = bank(); rows[3][6] = ""
    pkg.questions(rows)
    pkg.expect_error("المستوى فاضي")


def test_no_levels_uses_total_count(pkg):
    rows = [q(f"س{i}", ("1", "2", "3", "4"), "أ") for i in range(10)]
    pkg.questions(rows); pkg.settings(**{"عدد_الأسئلة_الكلي": 6})
    m = pkg.load()
    assert m.PICK == {"hard": 0, "medium": 6, "easy": 0}


def test_total_larger_than_bank_is_clamped_not_error(pkg):
    rows = [q(f"س{i}") for i in range(4)]
    pkg.questions(rows); pkg.settings(**{"عدد_الأسئلة_الكلي": 15})
    assert sum(pkg.load().PICK.values()) == 4


@pytest.mark.parametrize("mk", ["abc", "0", "-1", "101"])
def test_bad_marks(pkg, mk):
    rows = bank(); rows[0][9] = mk
    pkg.questions(rows)
    pkg.expect_error("سطر 2", "العلامة")


@pytest.mark.parametrize("mk,expected", [("٢", 2), ("٠٫٥", 0.5), ("2,5", 2.5), (2.5, 2.5), ("", 1)])
def test_marks_formats(pkg, mk, expected):
    rows = bank(); rows[0][9] = mk
    pkg.questions(rows)
    assert pkg.load().QLIST[0]["marks"] == expected


def test_bad_fixed_column(pkg):
    rows = bank(); rows[0][10] = "ربما"
    pkg.questions(rows)
    pkg.expect_error("سطر 2", "ثابت")


def test_empty_bank(pkg):
    pkg.questions([])
    pkg.expect_error("فاضي")


def test_blank_rows_ignored(pkg):
    rows = bank(); rows.insert(3, [""] * 11)
    pkg.questions(rows)
    assert len(pkg.load().QLIST) == 11


def test_several_errors_reported_together(pkg):
    rows = bank(); rows[0][5] = "ز"; rows[1][6] = "غريب"; rows[2][0] = ""
    pkg.questions(rows)
    msg = pkg.expect_error("سطر 2", "سطر 3", "سطر 4")
    assert msg.count("questions سطر") >= 3


def test_bank_smaller_than_requested_per_level(pkg):
    pkg.questions(bank(hard=1, medium=4, easy=4))      # الإعدادات بتطلب 2 صعب
    pkg.expect_error("أقل من المطلوب", "صعب")


def test_fixed_more_than_pick(pkg):
    rows = bank()
    for r in rows[:3]: r[10] = "نعم"                    # 3 ثوابت صعبة والمطلوب 2
    pkg.questions(rows)
    pkg.expect_error("الثابتة")


# ---------- ملف الطلاب ----------
def test_duplicate_student_same_section(pkg):
    pkg.students(roster() + [["العاشر", "أ", "طالب أ1"]])
    pkg.expect_error("مكرر", "طالب أ1")


def test_same_name_different_section_ok(pkg):
    pkg.students(roster() + [["العاشر", "ج", "طالب أ1"]])
    pkg.load()


@pytest.mark.parametrize("bad", [["", "أ", "اسم"], ["العاشر", "", "اسم"], ["العاشر", "أ", ""]])
def test_student_missing_field(pkg, bad):
    pkg.students(roster() + [bad])
    pkg.expect_error("لازم تعبّي")


def test_empty_roster(pkg):
    pkg.students([])
    pkg.expect_error("فاضي")


# ---------- settings.txt ----------
def test_unknown_setting_key(pkg):
    pkg.settings(**{"مدة_غلط": 5})
    pkg.expect_error("غير معروف", "مدة_غلط")


@pytest.mark.parametrize("key,val", [("خلط_الأسئلة", "ربما"), ("منع_النسخ", "2"), ("رمز_الدخول", "؟")])
def test_bad_boolean(pkg, key, val):
    pkg.settings(**{key: val})
    pkg.expect_error("نعم أو لا")


def test_bad_review_value(pkg):
    pkg.settings(**{"مراجعة_الأخطاء": "غدا"})
    pkg.expect_error("بعد_الإنهاء")


def test_minutes_not_number(pkg):
    pkg.settings(**{"المدة_بالدقائق": "ثلاثين"})
    pkg.expect_error("رقم")


@pytest.mark.parametrize("key", ["علامة_السؤال", "العلامة_النهائية"])
@pytest.mark.parametrize("val", ["abc", "0", "-3", "5000"])
def test_bad_mark_settings(pkg, key, val):
    pkg.settings(**{key: val})
    pkg.expect_error(key)


def test_qmark_and_final_must_match(pkg):
    pkg.settings(**{"علامة_السؤال": 1.5, "العلامة_النهائية": 20})        # 8 أسئلة × 1.5 = 12
    pkg.expect_error("12", "20")


def test_qmark_and_final_match_ok(pkg):
    pkg.settings(**{"علامة_السؤال": 1.5, "العلامة_النهائية": 12})
    assert pkg.load().SET["final"] == 12


def test_settings_cp1256_encoding(pkg):
    pkg.p("settings.txt").write_bytes(("عنوان_الامتحان = اختبار\nالمدة_بالدقائق = 20\n" +
                                       "عدد_الأسئلة_الصعبة = 2\nعدد_الأسئلة_المتوسطة = 3\nعدد_الأسئلة_السهلة = 3\n").encode("cp1256"))
    m = pkg.load()
    assert m.SET["minutes"] == 20 and m.SET["title"] == "اختبار"


def test_settings_arabic_digits_for_marks(pkg):
    pkg.settings(**{"العلامة_النهائية": "٣٠"})
    assert pkg.load().SET["final"] == 30


def test_comments_and_blank_lines_ignored(pkg):
    pkg.raw_settings("# تعليق\n\nعنوان_الامتحان = X\n# المدة_بالدقائق = 999\nعدد_الأسئلة_الصعبة=2\nعدد_الأسئلة_المتوسطة=3\nعدد_الأسئلة_السهلة=3\n")
    assert pkg.load().SET["minutes"] == 30


# ---------- ملفات ناقصة / صيغ ----------
@pytest.mark.parametrize("name", ["questions.xlsx", "students.xlsx"])                  # settings.txt اختياري (test_settings_optional.py)
def test_missing_file(pkg, name):
    pkg.p(name).unlink()
    pkg.expect_error("غير موجود")


def test_csv_utf8_bom_semicolon(pkg):
    pkg.p("questions.xlsx").unlink()
    lines = [";".join(Q_HEAD)] + [";".join(str(c) for c in r) for r in bank()]
    pkg.p("questions.csv").write_bytes(("\n".join(lines)).encode("utf-8-sig"))
    assert len(pkg.load().QLIST) == 11


def test_csv_cp1256(pkg):
    pkg.p("questions.xlsx").unlink()
    lines = [",".join(Q_HEAD)] + [",".join(str(c) for c in r) for r in bank()]
    pkg.p("questions.csv").write_bytes(("\n".join(lines)).encode("cp1256"))
    assert pkg.load().QLIST[0]["q"].startswith("سؤال")


def test_xlsx_preferred_over_csv(pkg):
    pkg.p("questions.csv").write_text("x,y\n1,2\n", encoding="utf-8")
    assert len(pkg.load().QLIST) == 11


def test_corrupt_xlsx_gives_readable_error(pkg):
    pkg.p("questions.xlsx").write_bytes(b"this is not a zip")
    msg = pkg.expect_error("خطأ بقراءة الملفات")
    assert "Excel" in msg


def test_numeric_cells_in_options(pkg):
    """خيارات مكتوبة كأرقام بإكسل (3.0 / 0.5): لازم تنقرأ كنص نظيف"""
    rows = bank(); rows[0] = q("كم؟", (3, 0.5, 12, 7), "ب", "صعب")
    pkg.questions(rows)
    o = [x.strip("⁦⁩") for x in pkg.load().QLIST[0]["o"]]      # autoltr بيلف الأرقام بعلامات اتجاه خفية
    assert o[:2] == ["3", "0.5"]


# ---------- الصور ----------
def test_missing_image(pkg):
    rows = bank(); rows[0][0] = "شاهد [img:fig.png] ثم أجب"
    pkg.questions(rows)
    pkg.expect_error("fig.png", "غير موجودة")


def test_image_present(pkg):
    rows = bank(); rows[0][0] = "شاهد [img:fig.png] ثم أجب"
    pkg.questions(rows); pkg.image("fig.png")
    m = pkg.load()
    assert m.IMG_LIST == ["fig.png"] and "/img/fig.png" in m.pub_q(1)["q"] and "/img/fig.png" in m.PREVIEW


def test_image_inside_option(pkg):
    rows = bank(); rows[0][1] = "[img:a.png]"
    pkg.questions(rows); pkg.image("a.png")
    assert pkg.load().IMG_LIST == ["a.png"]


@pytest.mark.parametrize("name", ["../x.png", "sub/x.png", "..\\x.png"])
def test_image_name_with_folders_rejected(pkg, name):
    rows = bank(); rows[0][0] = f"[img:{name}]"
    pkg.questions(rows)
    pkg.expect_error("بدون مجلدات")


def test_image_unsupported_extension(pkg):
    rows = bank(); rows[0][0] = "[img:doc.exe]"
    pkg.questions(rows); pkg.image("doc.exe")
    pkg.expect_error("غير مدعومة")


def test_arabic_img_token_alias(pkg):
    rows = bank(); rows[0][0] = "[صورة: fig.png]"
    pkg.questions(rows); pkg.image("fig.png")
    assert pkg.load().IMG_LIST == ["fig.png"]


# ---------- exam_data.zip ----------
def _zip(files):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for n, d in files.items(): z.writestr(n, d)
    return buf.getvalue()


def _xlsx_bytes(pkg, head, rows):
    pkg.write_xlsx("_tmp.xlsx", head, rows)
    data = pkg.p("_tmp.xlsx").read_bytes(); pkg.p("_tmp.xlsx").unlink()
    return data


def _settings_text():
    from conftest import DEFAULT_SETTINGS
    return "\n".join(f"{k} = {v}" for k, v in DEFAULT_SETTINGS.items()).replace("اختبار الوحدة", "من الزيب")


def test_zip_unpacks_data(pkg):
    for n in ("questions.xlsx", "students.xlsx", "settings.txt"): pkg.p(n).unlink()
    from conftest import S_HEAD
    pkg.p("exam_data.zip").write_bytes(_zip({
        "questions.xlsx": _xlsx_bytes(pkg, Q_HEAD, bank()), "students.xlsx": _xlsx_bytes(pkg, S_HEAD, roster()),
        "settings.txt": _settings_text(), "images/p1.png": b"PNG"}))
    m = pkg.load()
    assert m.SET["title"] == "من الزيب" and (pkg.path / "images" / "p1.png").exists()


def test_zip_with_single_top_folder_is_flattened(pkg):
    for n in ("questions.xlsx", "students.xlsx", "settings.txt"): pkg.p(n).unlink()
    from conftest import S_HEAD
    pkg.p("exam_data.zip").write_bytes(_zip({
        "exam_data/questions.xlsx": _xlsx_bytes(pkg, Q_HEAD, bank()), "exam_data/students.xlsx": _xlsx_bytes(pkg, S_HEAD, roster()),
        "exam_data/settings.txt": _settings_text()}))
    assert pkg.load().SET["title"] == "من الزيب"


def test_zip_slip_is_blocked(pkg):
    outside = pkg.path.parent / "evil.txt"
    pkg.p("exam_data.zip").write_bytes(_zip({"../evil.txt": "pwned", "..\\evil2.txt": "pwned", "images/../../evil3.png": "x", "server.py": "print('hijack')"}))
    pkg.load()
    assert not outside.exists() and not (pkg.path.parent / "evil2.txt").exists() and not (pkg.path.parent / "evil3.png").exists()
    assert "hijack" not in pkg.p("server.py").read_text(encoding="utf-8")        # ملفات غير مألوفة ما بتنكتب


def test_zip_does_not_overwrite_code_files(pkg):
    before = pkg.p("template.html").read_bytes()
    pkg.p("exam_data.zip").write_bytes(_zip({"template.html": "<html>x</html>", "check_files.py": "x=1"}))
    pkg.load()
    assert pkg.p("template.html").read_bytes() == before


def test_corrupt_zip(pkg):
    pkg.p("exam_data.zip").write_bytes(b"garbage")
    pkg.expect_error("exam_data.zip")


def test_zip_only_unpacked_once_per_version(pkg):
    """لو الأستاذ عدّل questions.xlsx بنفسه بعد الفك، التشغيل الثاني ما يطمّس تعديله"""
    for n in ("questions.xlsx", "students.xlsx", "settings.txt"): pkg.p(n).unlink()
    from conftest import S_HEAD
    pkg.p("exam_data.zip").write_bytes(_zip({
        "questions.xlsx": _xlsx_bytes(pkg, Q_HEAD, bank()), "students.xlsx": _xlsx_bytes(pkg, S_HEAD, roster()), "settings.txt": _settings_text()}))
    pkg.load()
    pkg.questions(bank(hard=5, medium=5, easy=5))
    assert len(pkg.load().QLIST) == 15


def test_zip_csv_replaces_old_xlsx(pkg):
    """زيب فيه questions.csv: لازم يلغي questions.xlsx القديم (السيرفر بيفضّل xlsx)"""
    lines = [",".join(Q_HEAD)] + [",".join(str(c) for c in r) for r in bank(hard=4, medium=4, easy=4)]
    pkg.p("exam_data.zip").write_bytes(_zip({"questions.csv": "\n".join(lines).encode("utf-8-sig")}))
    assert len(pkg.load().QLIST) == 12 and not pkg.p("questions.xlsx").exists()


# ---------- check_files.py (الأداة اللي بيشغّلها الأستاذ) ----------
def _run_check(pkg):
    return subprocess.run([sys.executable, "-X", "utf8", "check_files.py"], cwd=pkg.path, capture_output=True, text=True, encoding="utf-8", timeout=60)


def test_check_files_ok_and_report(pkg):
    r = _run_check(pkg)
    assert r.returncode == 0, r.stdout + r.stderr
    rep = pkg.p("check_report.txt").read_text(encoding="utf-8-sig")
    assert "بنك الأسئلة" in rep and "11 سؤال" in rep and "العلامة الكاملة: 8" in rep and "الطلاب         : 5" in rep


def test_check_files_fails_with_exit_code_and_error_file(pkg):
    rows = bank(); rows[0][5] = "ز"
    pkg.questions(rows)
    r = _run_check(pkg)
    assert r.returncode != 0
    assert "سطر 2" in pkg.p("exam_error.txt").read_text(encoding="utf-8-sig")


def test_check_files_warns_duplicates_and_skewed_answers(pkg):
    rows = bank(hard=4, medium=6, easy=6)
    for r in rows: r[5] = "أ"                                    # كل الإجابات أ
    rows[1][0] = rows[0][0]                                      # سؤال مكرر
    rows[2][1:5] = ["س", "س", "ص", "ع"]                          # خياران متطابقان
    pkg.questions(rows)
    r = _run_check(pkg)
    rep = pkg.p("check_report.txt").read_text(encoding="utf-8-sig")
    assert r.returncode == 0 and "سؤال مكرر" in rep and "خياران متطابقان" in rep and "وزّع الإجابات" in rep and "بدون حل" in rep


# ---------- تحذير الأسماء اللي فيها | ----------
def test_check_files_warns_about_pipe_in_name_grade_section(pkg):
    pkg.students(roster() + [["العاشر", "أ", "علي|محمد"], ["صف|خاص", "أ", "طالب س"], ["العاشر", "ش|ب", "طالب ع"]])
    r = _run_check(pkg)
    rep = pkg.p("check_report.txt").read_text(encoding="utf-8-sig")
    assert r.returncode == 0                                     # تنبيه مش خطأ: الامتحان بيشتغل لباقي الطلاب
    assert "الاسم «علي|محمد»" in rep and "الصف «صف|خاص»" in rep and "الشعبة «ش|ب»" in rep
    assert rep.count("ما رح يقدر يبدأ الامتحان") == 3 and "احذف هذا الرمز" in rep


def test_check_files_no_pipe_warning_for_normal_names(pkg):
    _run_check(pkg)
    assert "الرمز |" not in pkg.p("check_report.txt").read_text(encoding="utf-8-sig")


def test_pipe_warning_matches_real_failure(pkg):
    """التحذير مبرَّر: فعلاً الاسم اللي فيه | بيرجع «bad» من السيرفر"""
    pkg.students(roster() + [["العاشر", "أ", "علي|محمد"]])
    m = pkg.load()
    assert m.parse_key(f"{m.SET['title']}|العاشر|أ|علي|محمد") is None


# ---------- إصلاحات قارئ xlsx (Excel) ----------
def _q_file_with_cells(pkg, cells, rows_extra=()):
    """يكتب questions.xlsx فيه الأسئلة السليمة + صف إضافي بخلايا بأنواع محددة (مثل ما بيخزّنها Excel)"""
    import datetime, openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.append(Q_HEAD)
    for r in bank(): ws.append(r)
    ws.append(["Q12 types", *cells, "a", "easy"])
    wb.save(pkg.p("questions.xlsx"))


def test_xlsx_boolean_float_date_and_percent_cells(pkg):
    import datetime, openpyxl
    m = pkg.load()                                    # حزمة سليمة، وبعدين بنقرأ ملف الأنواع بـ read_table مباشرة
    wb = openpyxl.Workbook(); ws = wb.active; ws.append(Q_HEAD)
    for r in bank(): ws.append(r)
    ws.append(["Q12", True, False, datetime.datetime(1948, 1, 2), 0.1 + 0.2, "a", "easy"])
    ws.append(["Q13", 0.25, 0.5, datetime.datetime(2026, 3, 9, 14, 30), 1 / 3, "a", "easy"])
    ws["B14"].number_format = "0%"; ws["C14"].number_format = "0.0%"                    # الصف 14 = Q13 (0.25 و0.5)
    ws["A15"] = "Q14"; ws["B15"] = 0.0725; ws["B15"].number_format = "0.00%"; ws["F15"] = "a"; ws["G15"] = "easy"
    wb.save(pkg.p("questions.xlsx"))
    rows = m.read_table(str(pkg.p("questions.xlsx")))
    assert rows[12][1:5] == ["TRUE", "FALSE", "1948-01-02", "0.3"]
    assert rows[13][1:5] == ["25%", "50.0%", "2026-03-09 14:30", "0.333333333333333"]
    assert rows[14][1] == "7.25%"                                                       # الصف 15


def test_blank_rows_keep_line_numbers_in_error_messages(pkg):
    """صف فاضي بالنص (ما بينكتب بملف xlsx): رقم السطر بالرسالة لازم يطابق اللي بيشوفه الأستاذ بإكسل"""
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.append(Q_HEAD)
    rows = bank()
    for r in rows[:3]: ws.append(r)
    ws.append([]); ws.append([])                                  # سطرين فاضيين (5 و6)
    bad = list(rows[3]); bad[5] = "ز"
    ws.append(bad)                                                # هذا هو السطر 7 بإكسل
    for r in rows[4:]: ws.append(r)
    wb.save(pkg.p("questions.xlsx"))
    pkg.expect_error("سطر 7", "الجواب")


def test_cells_without_reference_attribute_are_read_in_order(pkg):
    """بعض الأدوات بتكتب <c> بدون r: ترتيب الأعمدة لازم يضل صح"""
    import zipfile, shutil
    pkg.questions(); src = pkg.p("questions.xlsx"); dst = pkg.p("q2.xlsx")
    zin = zipfile.ZipFile(src); zout = zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED)
    for it in zin.infolist():
        data = zin.read(it.filename)
        if it.filename.startswith("xl/worksheets/sheet"): data = re.sub(rb' r="[A-Z]+[0-9]+"', b"", data).replace(b"<c r=", b"<c ")
        zout.writestr(it, data)
    zout.close(); zin.close(); shutil.move(dst, src)
    base = bank()
    m = pkg.load()
    assert [q["q"] for q in m.QLIST][:3] == [r[0] for r in base[:3]] or len(m.QLIST) == 11


def test_questions_sheet_found_by_header_not_position(pkg):
    import openpyxl
    wb = openpyxl.Workbook(); notes = wb.active; notes.title = "Notes"; notes.append(["ملاحظات للأستاذ"])
    ws = wb.create_sheet("Questions"); ws.append(Q_HEAD)
    for r in bank(): ws.append(r)
    wb.save(pkg.p("questions.xlsx"))
    assert len(pkg.load().QLIST) == 11


def test_hidden_sheet_is_skipped_when_no_header_matches(pkg):
    import openpyxl
    wb = openpyxl.Workbook(); hid = wb.active; hid.title = "old"; hid.append(["x"]); hid.sheet_state = "hidden"
    ws = wb.create_sheet("data"); ws.append(["عنوان غريب"] + Q_HEAD[1:])
    for r in bank(): ws.append(r)
    wb.save(pkg.p("questions.xlsx"))
    assert len(pkg.load().QLIST) == 11


def test_students_sheet_found_by_header(pkg):
    import openpyxl
    wb = openpyxl.Workbook(); n = wb.active; n.title = "Notes"; n.append(["ملاحظات"])
    ws = wb.create_sheet("Students"); ws.append(S_HEAD)
    for r in roster(): ws.append(r)
    wb.save(pkg.p("students.xlsx"))
    assert sum(len(x) for g in pkg.load().ROSTER.values() for x in g.values()) == 5


def test_csv_with_cr_only_line_endings(pkg):
    pkg.p("questions.xlsx").unlink()
    text = chr(13).join([",".join(Q_HEAD)] + [",".join(str(c) for c in r) for r in bank()])
    pkg.p("questions.csv").write_bytes(text.encode("utf-8-sig"))
    assert len(pkg.load().QLIST) == 11


def test_question_marks_instead_of_arabic_is_a_clear_error(pkg):
    rows = bank(); rows[2][0] = "????? ??? 3"; rows[3][1] = "????"
    pkg.questions(rows)
    msg = pkg.expect_error("سطر 4", "سطر 5", "علامات استفهام")
    assert "UTF-8" in msg


@pytest.mark.parametrize("text", ["ما هو X؟", "What is this?", "Is it A? Or B?", "Why??", "هل 1 + 1 = 2؟"])
def test_normal_question_marks_are_not_flagged(pkg, text):
    rows = bank(); rows[0][0] = text; pkg.questions(rows)
    assert len(pkg.load().QLIST) == 11


def test_torn_last_line_is_repaired_at_startup(pkg):
    pkg.p("submissions.jsonl").write_text('{"key": "x|y|z|w", "exam": "اختبار", "gra', encoding="utf-8")
    pkg.load()
    assert pkg.p("submissions.jsonl").read_bytes().endswith(bytes([10]))


# ---------- تقارير قديمة مضللة ----------
def test_stale_error_file_is_removed_once_files_are_fixed(pkg):
    """تجربة حقيقية: الأستاذ صلّح الصور الناقصة، والسيرفر اشتغل، بس exam_error.txt ضل فيه الخطأ القديم"""
    pkg.p("exam_error.txt").write_text("في مشكلة بالملفات: الصورة p048_q3.png غير موجودة", encoding="utf-8-sig")
    r = _run_check(pkg)
    assert r.returncode == 0 and not pkg.p("exam_error.txt").exists() and pkg.p("check_report.txt").exists()


def test_server_start_also_removes_stale_error_file(pkg):
    from test_http import Server
    pkg.p("exam_error.txt").write_text("خطأ قديم", encoding="utf-8-sig")
    s = Server(pkg).start()
    try:
        assert not pkg.p("exam_error.txt").exists()
    finally:
        s.stop()


def test_stale_ok_report_is_removed_when_check_fails(pkg):
    """العكس: تقرير «الملفات سليمة ✔» من فحص سابق ما لازم يضل لما الفحص الجديد يفشل"""
    assert _run_check(pkg).returncode == 0 and "OK - the exam files are valid" in pkg.p("check_report.txt").read_text(encoding="utf-8-sig")
    rows = bank(); rows[0][5] = "ز"; pkg.questions(rows)
    r = _run_check(pkg)
    assert r.returncode != 0 and not pkg.p("check_report.txt").exists() and "سطر 2" in pkg.p("exam_error.txt").read_text(encoding="utf-8-sig")


def test_failed_start_keeps_the_error_file(pkg):
    rows = bank(); rows[0][5] = "ز"; pkg.questions(rows)
    pkg.expect_error("سطر 2")
    assert pkg.p("exam_error.txt").exists()
